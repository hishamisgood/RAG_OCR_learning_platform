"""Ingestion + retrieval on a local ChromaDB store (heavy imports are lazy).

Ingestion is incremental: a small manifest remembers which files were indexed,
so adding one book does not re-embed all the others.

Scanned PDF pages and images are OCR'd (see ocr.py); OCR results are cached on
disk, so re-running ingest never pays for the same page twice.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Callable

from langchain_core.documents import Document

from .config import Settings

TEXT_EXT = {".md", ".txt"}
IMAGE_EXT = {".jpg", ".jpeg", ".png"}
SUPPORTED = TEXT_EXT | IMAGE_EXT | {".pdf", ".docx"}
MANIFEST = "manifest.json"

Progress = Callable[[str], None]


def _embeddings(s: Settings):
    from langchain_huggingface import HuggingFaceEmbeddings

    try:
        import torch

        device = "cuda" if torch.cuda.is_available() else "cpu"
    except ImportError:
        device = "cpu"
    return HuggingFaceEmbeddings(
        model_name=s.embedding_model,
        model_kwargs={"device": device},
        encode_kwargs={"normalize_embeddings": True, "batch_size": 16},
    )


_STORES: dict[tuple, object] = {}


def _store(s: Settings, fresh: bool = False):
    """One Chroma handle per (dir, collection, model) - loading bge-m3 takes a few seconds."""
    from langchain_chroma import Chroma

    key = (str(s.chroma_dir), s.collection, s.embedding_model)
    if fresh or key not in _STORES:
        _STORES[key] = Chroma(
            collection_name=s.collection,
            embedding_function=_embeddings(s) if key not in _STORES else _STORES[key]._embedding_function,
            persist_directory=str(s.chroma_dir),
        )
    return _STORES[key]


# ---- file discovery / planning (pure functions, easy to test) ---------------

def list_files(s: Settings) -> list[Path]:
    return [p for p in sorted(s.docs_dir.rglob("*")) if p.is_file() and p.suffix.lower() in SUPPORTED]


def _rel(s: Settings, p: Path) -> str:
    return p.relative_to(s.docs_dir).as_posix()


def subject_of(rel: str) -> str:
    """Folder for files in sub-folders ('Student/part 2.pdf' -> 'Student'), else the file name."""
    parts = rel.split("/")
    return "/".join(parts[:-1]) if len(parts) > 1 else Path(rel).stem


def _signature(p: Path) -> str:
    st = p.stat()
    return f"{st.st_size}:{int(st.st_mtime)}"


def plan_ingest(current: dict[str, str], manifest: dict[str, str]) -> tuple[list[str], list[str]]:
    """Compare files on disk with what was indexed. Returns (paths_to_remove, paths_to_add)."""
    changed = [k for k, sig in current.items() if manifest.get(k) != sig]
    removed = [k for k in manifest if k not in current]
    to_remove = removed + [k for k in changed if k in manifest]
    return to_remove, changed


def _read_manifest(s: Settings) -> dict[str, str] | None:
    f = s.chroma_dir / MANIFEST
    if not f.exists():
        return None
    try:
        return json.loads(f.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None


def _write_manifest(s: Settings, manifest: dict[str, str]) -> None:
    s.chroma_dir.mkdir(parents=True, exist_ok=True)
    (s.chroma_dir / MANIFEST).write_text(json.dumps(manifest, indent=1, ensure_ascii=False), encoding="utf-8")


# ---- loading ----------------------------------------------------------------

class LoadResult(list):
    """list[Document] plus bookkeeping: pages that needed OCR but could not get it."""

    failed_pages: int = 0


def load_file(s: Settings, path: Path, ocr=None, log: Progress = print) -> LoadResult:
    """Read one file into Documents (one per PDF page / image).

    `ocr` is an ocr.OCR instance (or None to skip scanned pages).
    """
    rel, suffix = _rel(s, path), path.suffix.lower()
    meta = {"source": path.name, "path": rel, "subject": subject_of(rel), "origin": "notes"}
    out = LoadResult()

    if suffix in TEXT_EXT:
        out.append(Document(page_content=path.read_text(encoding="utf-8"), metadata={**meta, "ocr": False}))
        return out

    if suffix == ".docx":
        import docx

        d = docx.Document(str(path))
        lines = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            lines += [" | ".join(c.text.strip() for c in row.cells) for row in t.rows]
        if lines:
            out.append(Document(page_content="\n".join(lines), metadata={**meta, "ocr": False}))
        return out

    from .ocr import file_hash, load_image, render_pdf_page

    if suffix in IMAGE_EXT:
        if ocr is None:
            out.failed_pages = 1
            return out
        text = ocr.page_text(file_hash(path), 1, lambda: (load_image(path), "image/jpeg"))
        if text.strip() and text.strip() != "[blank page]":
            out.append(Document(page_content=text, metadata={**meta, "page": 1, "ocr": True}))
        return out

    # PDF: use the text layer when it exists, OCR the page otherwise.
    import pymupdf

    fhash = None
    with pymupdf.open(str(path)) as doc:
        n = len(doc)
        texts = [p.get_text().strip() for p in doc]
        # A digital book (most pages have a real text layer) is not a scan: its figure and
        # chapter-divider pages must not be sent to OCR, or they burn the free quota.
        mostly_digital = n >= 20 and sum(len(t) >= s.ocr_min_chars for t in texts) / n >= 0.6
        if mostly_digital:
            log(f"    digital PDF ({n} pages): OCR skipped for figure/divider pages")
        for i, page in enumerate(doc):
            text = texts[i]
            if len(text) >= s.ocr_min_chars:
                out.append(Document(page_content=text, metadata={**meta, "page": i + 1, "ocr": False}))
                continue
            if mostly_digital:
                if text:
                    out.append(Document(page_content=text, metadata={**meta, "page": i + 1, "ocr": False}))
                continue
            if not page.get_images():  # blank / separator page - nothing to read
                continue
            if ocr is None:
                out.failed_pages += 1
                continue
            fhash = fhash or file_hash(path)
            hit = ocr.cached(fhash, i + 1) is not None
            if not hit:
                log(f"    OCR page {i + 1}/{n} of {path.name}")
            try:
                text = ocr.page_text(fhash, i + 1, lambda i=i: (render_pdf_page(doc, i), "image/png"))
            except Exception as e:
                if type(e).__name__ == "QuotaExhausted":
                    raise
                log(f"    [ocr failed] page {i + 1}: {str(e)[:120]}")
                out.failed_pages += 1
                continue
            if text.strip() and text.strip() != "[blank page]":
                out.append(Document(page_content=text, metadata={**meta, "page": i + 1, "ocr": True}))
    return out


def load_documents(s: Settings) -> list[Document]:
    """Read every supported file under data/docs without OCR (kept for backwards compatibility)."""
    docs: list[Document] = []
    for p in list_files(s):
        if p.stat().st_size:
            docs.extend(load_file(s, p))
    return docs


# ---- ingestion --------------------------------------------------------------

def ingest(s: Settings, rebuild: bool = False, log: Progress = print,
           on_file: Callable[[int, int, str], None] | None = None) -> int:
    """Index new/changed files, drop removed ones. Returns chunks added in this run.

    rebuild=True wipes the index and re-embeds everything (needed after changing chunk
    size or embedding model). OCR results are cached, so a rebuild does not re-OCR.
    A file with pages whose OCR failed is indexed but not marked done, so the next
    ingest retries just those pages.
    """
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    from .ocr import QuotaExhausted, make_ocr

    files = {_rel(s, p): p for p in list_files(s)}
    if not files:
        raise SystemExit(f"No supported files found in {s.docs_dir}")

    manifest = _read_manifest(s)
    if manifest is None:  # first run with this version (or no index yet): start clean
        rebuild = True
    store = _store(s)
    if rebuild:
        store.delete_collection()
        store = _store(s, fresh=True)
        manifest = {}

    current = {}
    for rel, p in files.items():
        if p.stat().st_size == 0:
            log(f"[skip] {rel}: file is empty (0 KB) - copy it again")
            continue
        current[rel] = _signature(p)
    to_remove, to_add = plan_ingest(current, manifest)

    for rel in to_remove:
        store._collection.delete(where={"path": rel})
        manifest.pop(rel, None)
        log(f"[removed] {rel}")

    ocr = make_ocr(s)
    splitter = RecursiveCharacterTextSplitter(chunk_size=s.chunk_size, chunk_overlap=s.chunk_overlap)
    total = 0
    for n, rel in enumerate(to_add, 1):
        if on_file:
            on_file(n, len(to_add), rel)
        log(f"[{n}/{len(to_add)}] {rel}")
        try:
            docs = load_file(s, files[rel], ocr, log)
        except QuotaExhausted as e:
            log(f"[stopped] {e}")
            break
        except Exception as e:  # a broken file should not stop the whole run
            log(f"[skip] {rel}: could not read ({e})")
            continue
        store._collection.delete(where={"path": rel})  # drop chunks from an earlier partial run
        chunks = splitter.split_documents(docs)
        for i in range(0, len(chunks), 64):
            store.add_documents(chunks[i:i + 64])
            log(f"    embedded {min(i + 64, len(chunks))}/{len(chunks)} chunks")
        total += len(chunks)
        ocr_pages = sum(1 for d in docs if d.metadata.get("ocr"))
        note = f", {ocr_pages} OCR pages" if ocr_pages else ""
        if docs.failed_pages and s.ocr_provider != "none":
            log(f"    {len(chunks)} chunks{note}; {docs.failed_pages} pages still need OCR (will retry next ingest)")
        else:
            manifest[rel] = current[rel]
            if docs.failed_pages:
                log(f"    {len(chunks)} chunks{note}; {docs.failed_pages} scanned pages skipped because OCR is disabled")
            else:
                log(f"    {len(chunks)} chunks{note}" if chunks else "    no text found")
        _write_manifest(s, manifest)  # save progress after each file
    _write_manifest(s, manifest)
    if ocr is not None and any(ocr.stats.values()):
        log(f"OCR pages: {ocr.stats}")
    if not to_add and not to_remove:
        log("Nothing new to index.")
    return total


# ---- retrieval / library ------------------------------------------------------

class Retriever:
    """`.invoke(query, subjects=None)` -> top-k chunks, optionally limited to some subjects."""

    def __init__(self, s: Settings):
        self.s = s

    def invoke(self, query: str, subjects: list[str] | None = None) -> list[Document]:
        flt = {"subject": {"$in": list(subjects)}} if subjects else None
        return _store(self.s).similarity_search(query, k=self.s.top_k, filter=flt)


def get_retriever(s: Settings) -> Retriever:
    return Retriever(s)


def library(s: Settings) -> list[dict]:
    """One row per indexed file: subject, chunks, pages, OCR pages."""
    # Listing the existing collection needs only its metadata.  Going through _store()
    # unnecessarily constructs the embedding model, which can make the UI wait for a
    # Hugging Face download before it can even show the library status.
    import chromadb

    client = chromadb.PersistentClient(path=str(s.chroma_dir))
    data = client.get_collection(s.collection).get(include=["metadatas"])
    rows: dict[str, dict] = {}
    for m in data.get("metadatas") or []:
        r = rows.setdefault(m.get("path", "?"), {"file": m.get("path", "?"), "subject": m.get("subject", ""),
                                                 "chunks": 0, "_pages": set(), "_ocr": set()})
        r["chunks"] += 1
        if "page" in m:
            r["_pages"].add(m["page"])
            if m.get("ocr"):
                r["_ocr"].add(m["page"])
    out = []
    for r in rows.values():
        pages, ocr = r.pop("_pages"), r.pop("_ocr")
        out.append({**r, "pages": len(pages), "ocr_pages": len(ocr)})
    return sorted(out, key=lambda r: r["file"])


def subjects(s: Settings) -> list[str]:
    return sorted({r["subject"] for r in library(s) if r["subject"]})